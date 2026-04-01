from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import datetime
import json
import os
import re
import shutil
import zipfile
from typing import Iterable

from config import ARTIFACTS_DIR

DEFAULT_WORKSPACE_TEMPLATE_ROOT = os.getenv(
    "WORKSPACE_TEMPLATE_ROOT",
    r"C:\Users\adamo\Dropbox\_Consulting\Organized Files\Templates",
)
WORKSPACE_TEMPLATE_CACHE_DIR = os.path.join(str(ARTIFACTS_DIR), "workspace_templates", "imported")

_MACHINE_TAG_RE = re.compile(r"\[\[([A-Z]+):([^\]]+)\]\]")
_BRACKET_ONLY_RE = re.compile(r"^\[[^\]]+\]$")
_WHITESPACE_RE = re.compile(r"\s+")


@dataclass(frozen=True)
class TemplateFillTarget:
    key: str
    token_kind: str
    token_name: str
    label: str
    order: int


@dataclass(frozen=True)
class WorkspaceTemplateSpec:
    key: str
    display_name: str
    title: str
    source: str
    machine_path: str
    human_path: str
    sections: list[str]
    instructions: list[str]
    fields: list[str]
    values: list[str]
    fill_targets: list[TemplateFillTarget]

    def to_payload(self) -> dict:
        return {
            "key": self.key,
            "display_name": self.display_name,
            "title": self.title,
            "source": self.source,
            "machine_path": self.machine_path,
            "human_path": self.human_path,
            "sections": list(self.sections),
            "instructions": list(self.instructions),
            "fields": list(self.fields),
            "values": list(self.values),
            "fill_targets": [asdict(t) for t in self.fill_targets],
        }

    @classmethod
    def from_payload(cls, payload: dict | None) -> "WorkspaceTemplateSpec | None":
        if not isinstance(payload, dict):
            return None
        fill_targets = []
        for raw in payload.get("fill_targets") or []:
            if not isinstance(raw, dict):
                continue
            fill_targets.append(
                TemplateFillTarget(
                    key=str(raw.get("key") or ""),
                    token_kind=str(raw.get("token_kind") or ""),
                    token_name=str(raw.get("token_name") or ""),
                    label=str(raw.get("label") or ""),
                    order=int(raw.get("order") or 0),
                )
            )
        return cls(
            key=str(payload.get("key") or ""),
            display_name=str(payload.get("display_name") or ""),
            title=str(payload.get("title") or ""),
            source=str(payload.get("source") or "external"),
            machine_path=str(payload.get("machine_path") or ""),
            human_path=str(payload.get("human_path") or ""),
            sections=[str(v) for v in (payload.get("sections") or [])],
            instructions=[str(v) for v in (payload.get("instructions") or [])],
            fields=[str(v) for v in (payload.get("fields") or [])],
            values=[str(v) for v in (payload.get("values") or [])],
            fill_targets=fill_targets,
        )


@dataclass(frozen=True)
class TemplateValidationResult:
    is_valid: bool
    missing_block_keys: list[str]
    empty_block_keys: list[str]
    unexpected_block_keys: list[str]
    missing_metadata_keys: list[str]
    unexpected_metadata_keys: list[str]

    def issues(self) -> list[str]:
        lines: list[str] = []
        if self.missing_block_keys:
            lines.append(f"Missing template block keys: {', '.join(self.missing_block_keys)}")
        if self.empty_block_keys:
            lines.append(f"Empty template block values: {', '.join(self.empty_block_keys)}")
        if self.unexpected_block_keys:
            lines.append(f"Unexpected template block keys: {', '.join(self.unexpected_block_keys)}")
        if self.missing_metadata_keys:
            lines.append(f"Missing document metadata keys: {', '.join(self.missing_metadata_keys)}")
        if self.unexpected_metadata_keys:
            lines.append(f"Unexpected document metadata keys: {', '.join(self.unexpected_metadata_keys)}")
        return lines


def _normalize_template_base_name(name: str) -> str:
    base = str(name or "").strip()
    patterns = (
        "_machine",
        "_human",
        " - Machine Readable",
        " - Human Readable",
        " - CF Template",
    )
    for suffix in patterns:
        if base.endswith(suffix):
            base = base[: -len(suffix)]
            break
    base = base.strip(" _-")
    return base


def _humanize_name(name: str) -> str:
    text = _normalize_template_base_name(os.path.splitext(os.path.basename(str(name or "")))[0])
    text = text.replace("_", " ")
    text = _WHITESPACE_RE.sub(" ", text).strip()
    if not text:
        return "Untitled Template"
    return " ".join(part[:1].upper() + part[1:] for part in text.split(" "))


def _iter_docx_xml_text(docx_path: str) -> Iterable[str]:
    if not docx_path or not os.path.exists(docx_path):
        return []
    texts: list[str] = []
    try:
        with zipfile.ZipFile(docx_path, "r") as archive:
            for name in archive.namelist():
                if not name.startswith("word/"):
                    continue
                if not name.lower().endswith(".xml"):
                    continue
                if not (
                    name.startswith("word/document")
                    or name.startswith("word/header")
                    or name.startswith("word/footer")
                ):
                    continue
                try:
                    texts.append(archive.read(name).decode("utf-8", "ignore"))
                except Exception:
                    continue
    except Exception:
        return []
    return texts


def _extract_machine_tokens(machine_path: str) -> list[tuple[str, str]]:
    tokens: list[tuple[str, str]] = []
    for xml_text in _iter_docx_xml_text(machine_path):
        for kind, name in _MACHINE_TAG_RE.findall(xml_text):
            tokens.append((str(kind).strip().upper(), str(name).strip()))
    return tokens


def _build_fill_targets(tokens: list[tuple[str, str]]) -> list[TemplateFillTarget]:
    out: list[TemplateFillTarget] = []
    last_field_name = ""
    order = 0
    for kind, name in tokens:
        if kind == "FIELD":
            last_field_name = name
            continue
        if kind == "INSTRUCTION":
            out.append(
                TemplateFillTarget(
                    key=f"instruction::{name}",
                    token_kind=kind,
                    token_name=name,
                    label=_humanize_name(name),
                    order=order,
                )
            )
            order += 1
            continue
        if kind == "VALUE":
            label_name = last_field_name or name
            out.append(
                TemplateFillTarget(
                    key=f"value::{name}",
                    token_kind=kind,
                    token_name=name,
                    label=_humanize_name(label_name),
                    order=order,
                )
            )
            order += 1
            last_field_name = ""
    return out


def build_template_contract(spec: WorkspaceTemplateSpec) -> str:
    schema = build_template_schema(spec)
    lines = [
        f"Selected document template: {spec.display_name}",
        f"Machine template path: {spec.machine_path}",
        f"Human template path: {spec.human_path}",
        "",
        "You must produce BOTH:",
        "1. A finished markdown document suitable for review in the Workspace pane.",
        "2. A template_blocks JSON object that fills every required document block.",
        "",
        "Return template_blocks using these exact keys:",
    ]
    for target in spec.fill_targets:
        lines.append(f'- "{target.key}" -> {target.label}')
    lines.extend(
        [
            "",
            "Also return document_metadata with these keys:",
            '- "document_title"',
            '- "subtitle"',
            '- "document_id"',
            '- "version"',
            '- "effective_date"',
            '- "prepared_by"',
            "",
            "Schema object to follow exactly:",
            json.dumps(schema, indent=2),
            "",
            "Rules for template_blocks:",
            "- Every listed key must be present.",
            "- Use plain strings for each value.",
            "- Do not leave placeholders or bracketed instructions in the content.",
            "- Keep content faithful to the provided source files and user goal.",
            "",
            "Also return completion-analysis fields:",
            '- "evidence_map": object keyed by template block key',
            '- "unresolved_fields": list of template block keys still not fully supported by source evidence',
            '- "research_gaps": list of public/best-practice gaps that could be filled by research',
            '- "user_questions": list of targeted user/client questions required to finish the document safely',
        ]
    )
    return "\n".join(lines)


def _template_source_dirs() -> list[tuple[str, str]]:
    dirs: list[tuple[str, str]] = []
    external_root = os.path.abspath(str(DEFAULT_WORKSPACE_TEMPLATE_ROOT))
    if external_root:
        dirs.append(("external", external_root))
    cache_root = os.path.abspath(str(WORKSPACE_TEMPLATE_CACHE_DIR))
    dirs.append(("imported", cache_root))
    return dirs


def discover_workspace_templates() -> list[WorkspaceTemplateSpec]:
    machine_candidates: dict[tuple[str, str], str] = {}
    human_candidates: dict[tuple[str, str], str] = {}

    for source, root in _template_source_dirs():
        if not os.path.isdir(root):
            continue
        for base_dir, _dirs, files in os.walk(root):
            for filename in files:
                lower = filename.lower()
                full_path = os.path.abspath(os.path.join(base_dir, filename))
                stem, ext = os.path.splitext(filename)
                norm = _normalize_template_base_name(stem)
                key = (source, norm.lower())
                if lower.endswith("_machine.docx") or lower.endswith(" - machine readable.docx"):
                    machine_candidates[key] = full_path
                    continue
                if (
                    lower.endswith("_human.docx")
                    or lower.endswith("_human.dotx")
                    or lower.endswith(" - human readable.docx")
                    or lower.endswith(" - cf template.dotx")
                ):
                    human_candidates[key] = full_path

    specs: list[WorkspaceTemplateSpec] = []
    for (source, normalized_name), machine_path in sorted(machine_candidates.items()):
        human_path = human_candidates.get((source, normalized_name))
        if not human_path:
            continue
        tokens = _extract_machine_tokens(machine_path)
        fill_targets = _build_fill_targets(tokens)
        specs.append(
            WorkspaceTemplateSpec(
                key=f"{source}:{normalized_name}",
                display_name=_humanize_name(normalized_name),
                title=_humanize_name(normalized_name),
                source=source,
                machine_path=machine_path,
                human_path=human_path,
                sections=[name for kind, name in tokens if kind == "SECTION"],
                instructions=[name for kind, name in tokens if kind == "INSTRUCTION"],
                fields=[name for kind, name in tokens if kind == "FIELD"],
                values=[name for kind, name in tokens if kind == "VALUE"],
                fill_targets=fill_targets,
            )
        )
    return specs


def get_workspace_template_by_key(template_key: str) -> WorkspaceTemplateSpec | None:
    desired = str(template_key or "").strip()
    if not desired:
        return None
    for spec in discover_workspace_templates():
        if spec.key == desired:
            return spec
    return None


def build_template_schema(spec: WorkspaceTemplateSpec) -> dict:
    required_metadata = [
        "document_title",
        "subtitle",
        "document_id",
        "version",
        "effective_date",
        "prepared_by",
    ]
    return {
        "template_key": spec.key,
        "template_name": spec.display_name,
        "required_template_blocks": [target.key for target in spec.fill_targets],
        "template_block_labels": {target.key: target.label for target in spec.fill_targets},
        "required_document_metadata": required_metadata,
        "completion_analysis_keys": [
            "evidence_map",
            "unresolved_fields",
            "research_gaps",
            "user_questions",
        ],
        "notes": [
            "Return every required template block key exactly once.",
            "Return plain strings only for template_blocks values.",
            "Do not emit extra template block keys.",
            "Leave unresolved template block values blank rather than guessing.",
        ],
    }


def validate_template_payload(
    spec: WorkspaceTemplateSpec,
    *,
    template_blocks: dict[str, str] | None,
    document_metadata: dict | None,
) -> TemplateValidationResult:
    blocks = {str(k): str(v or "") for k, v in dict(template_blocks or {}).items()}
    metadata = {str(k): str(v or "") for k, v in dict(document_metadata or {}).items()}
    schema = build_template_schema(spec)
    required_block_keys = list(schema.get("required_template_blocks") or [])
    required_metadata_keys = list(schema.get("required_document_metadata") or [])

    missing_block_keys = [key for key in required_block_keys if key not in blocks]
    empty_block_keys = [key for key in required_block_keys if key in blocks and not str(blocks.get(key) or "").strip()]
    unexpected_block_keys = [key for key in blocks.keys() if key not in required_block_keys]
    missing_metadata_keys = [key for key in required_metadata_keys if key not in metadata]
    unexpected_metadata_keys = [key for key in metadata.keys() if key not in required_metadata_keys]

    return TemplateValidationResult(
        is_valid=not any(
            [
                missing_block_keys,
                empty_block_keys,
                unexpected_block_keys,
                missing_metadata_keys,
                unexpected_metadata_keys,
            ]
        ),
        missing_block_keys=missing_block_keys,
        empty_block_keys=empty_block_keys,
        unexpected_block_keys=unexpected_block_keys,
        missing_metadata_keys=missing_metadata_keys,
        unexpected_metadata_keys=unexpected_metadata_keys,
    )


def import_workspace_template_pair(
    *,
    machine_path: str,
    human_path: str,
    display_name: str | None = None,
) -> WorkspaceTemplateSpec:
    machine_abs = os.path.abspath(str(machine_path or ""))
    human_abs = os.path.abspath(str(human_path or ""))
    if not os.path.exists(machine_abs):
        raise FileNotFoundError(machine_abs)
    if not os.path.exists(human_abs):
        raise FileNotFoundError(human_abs)

    base_name = _normalize_template_base_name(display_name or os.path.splitext(os.path.basename(machine_abs))[0])
    slug = re.sub(r"[^a-zA-Z0-9_ -]", "", base_name).strip().replace(" ", "_") or "imported_template"
    target_dir = os.path.join(str(WORKSPACE_TEMPLATE_CACHE_DIR), slug)
    os.makedirs(target_dir, exist_ok=True)

    canonical_base = _normalize_template_base_name(slug) or "imported_template"
    machine_target = os.path.join(target_dir, f"{canonical_base}_machine.docx")
    human_ext = os.path.splitext(human_abs)[1].lower() or ".docx"
    if human_ext not in {".docx", ".dotx"}:
        human_ext = ".docx"
    human_target = os.path.join(target_dir, f"{canonical_base}_human{human_ext}")
    shutil.copy2(machine_abs, machine_target)
    shutil.copy2(human_abs, human_target)

    tokens = _extract_machine_tokens(machine_target)
    fill_targets = _build_fill_targets(tokens)
    spec = WorkspaceTemplateSpec(
        key=f"imported:{_normalize_template_base_name(slug).lower()}",
        display_name=_humanize_name(base_name),
        title=_humanize_name(base_name),
        source="imported",
        machine_path=os.path.abspath(machine_target),
        human_path=os.path.abspath(human_target),
        sections=[name for kind, name in tokens if kind == "SECTION"],
        instructions=[name for kind, name in tokens if kind == "INSTRUCTION"],
        fields=[name for kind, name in tokens if kind == "FIELD"],
        values=[name for kind, name in tokens if kind == "VALUE"],
        fill_targets=fill_targets,
    )
    return get_workspace_template_by_key(spec.key) or spec


def _default_document_metadata(spec: WorkspaceTemplateSpec, document_metadata: dict | None) -> dict[str, str]:
    raw = dict(document_metadata or {})
    now = datetime.now()
    return {
        "document_title": str(raw.get("document_title") or spec.title or spec.display_name),
        "subtitle": str(raw.get("subtitle") or ""),
        "document_id": str(raw.get("document_id") or "DRAFT"),
        "version": str(raw.get("version") or "0.1"),
        "effective_date": str(raw.get("effective_date") or now.strftime("%Y-%m-%d")),
        "prepared_by": str(raw.get("prepared_by") or "NaviSsurance"),
    }


def build_template_preview_markdown(
    spec: WorkspaceTemplateSpec,
    *,
    template_blocks: dict[str, str] | None,
    document_metadata: dict | None = None,
) -> str:
    blocks = {str(k): str(v or "").strip() for k, v in dict(template_blocks or {}).items()}
    meta = _default_document_metadata(spec, document_metadata)
    lines = [f"# {meta['document_title']}"]
    if meta["subtitle"]:
        lines.extend(["", meta["subtitle"]])
    lines.extend(
        [
            "",
            f"Document ID: {meta['document_id']}",
            f"Version: {meta['version']}",
            f"Effective Date: {meta['effective_date']}",
            f"Prepared by: {meta['prepared_by']}",
            "",
        ]
    )
    for target in spec.fill_targets:
        value = blocks.get(target.key, "").strip()
        if not value:
            continue
        if target.token_kind == "INSTRUCTION":
            lines.extend([value, ""])
        else:
            lines.extend([f"## {target.label}", value, ""])
    return "\n".join(lines).strip()


def render_workspace_template_to_docx(
    *,
    spec: WorkspaceTemplateSpec,
    template_blocks: dict[str, str] | None,
    output_path: str,
    document_metadata: dict | None = None,
) -> str:
    from docx import Document  # type: ignore

    human_template_path = os.path.abspath(str(spec.human_path))
    if not os.path.exists(human_template_path):
        raise FileNotFoundError(human_template_path)

    blocks = {str(k): str(v or "").strip() for k, v in dict(template_blocks or {}).items()}
    metadata = _default_document_metadata(spec, document_metadata)
    os.makedirs(os.path.dirname(os.path.abspath(str(output_path))), exist_ok=True)

    doc = Document(human_template_path)
    ordered_replacements = [blocks.get(target.key, "").strip() for target in spec.fill_targets]
    fill_index = 0

    for paragraph in _iter_document_paragraphs(doc):
        original = paragraph.text or ""
        updated = original
        updated = updated.replace("[Document Title]", metadata["document_title"])
        updated = updated.replace("[Optional Subtitle or Description]", metadata["subtitle"])
        updated = updated.replace("[ID]", metadata["document_id"])
        updated = updated.replace("[Version]", metadata["version"])
        updated = updated.replace("[Date]", metadata["effective_date"])
        updated = updated.replace("[Name / Company]", metadata["prepared_by"])
        updated = updated.replace("[Start Document Content]", "")
        stripped = updated.strip()
        if _BRACKET_ONLY_RE.fullmatch(stripped or ""):
            replacement = ""
            if fill_index < len(ordered_replacements):
                replacement = ordered_replacements[fill_index]
                fill_index += 1
            updated = replacement
        if updated != original:
            paragraph.text = updated

    doc.save(os.path.abspath(str(output_path)))
    return os.path.abspath(str(output_path))


def _iter_document_paragraphs(doc) -> Iterable:
    for paragraph in getattr(doc, "paragraphs", []):
        yield paragraph
    for table in getattr(doc, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
