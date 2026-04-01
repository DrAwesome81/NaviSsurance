from __future__ import annotations

from pathlib import Path

import pytest

pytest.importorskip("docx")

from docx import Document

from core import workspace_templates as wt


def _write_docx(path: Path, paragraphs: list[str]) -> None:
    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    doc.save(str(path))


def test_discover_workspace_templates_extracts_machine_tokens(tmp_path, monkeypatch):
    external_root = tmp_path / "external"
    cache_root = tmp_path / "cache"
    external_root.mkdir()
    cache_root.mkdir()
    monkeypatch.setattr(wt, "DEFAULT_WORKSPACE_TEMPLATE_ROOT", str(external_root))
    monkeypatch.setattr(wt, "WORKSPACE_TEMPLATE_CACHE_DIR", str(cache_root))

    _write_docx(
        external_root / "risk_management_plan_machine.docx",
        [
            "[[SECTION:risk_management_plan]]",
            "[[INSTRUCTION:risk_management_plan_scope]]",
            "[[FIELD:product_scope]]",
            "[[VALUE:product_scope]]",
            "[[FIELD:risk_acceptability_criteria]]",
            "[[VALUE:risk_acceptability_criteria]]",
        ],
    )
    _write_docx(
        external_root / "risk_management_plan_human.docx",
        [
            "[Document Title]",
            "Document ID: [ID]",
            "Risk Management Plan",
            "[Describe the scope of the plan.]",
            "Product Scope",
            "[Describe the product scope.]",
            "Risk Acceptability Criteria",
            "[Describe the risk acceptability criteria.]",
        ],
    )

    specs = wt.discover_workspace_templates()

    assert len(specs) == 1
    spec = specs[0]
    assert spec.display_name == "Risk Management Plan"
    assert spec.sections == ["risk_management_plan"]
    assert spec.instructions == ["risk_management_plan_scope"]
    assert spec.values == ["product_scope", "risk_acceptability_criteria"]
    assert [target.key for target in spec.fill_targets] == [
        "instruction::risk_management_plan_scope",
        "value::product_scope",
        "value::risk_acceptability_criteria",
    ]


def test_render_workspace_template_to_docx_fills_human_template(tmp_path, monkeypatch):
    external_root = tmp_path / "external"
    cache_root = tmp_path / "cache"
    external_root.mkdir()
    cache_root.mkdir()
    monkeypatch.setattr(wt, "DEFAULT_WORKSPACE_TEMPLATE_ROOT", str(external_root))
    monkeypatch.setattr(wt, "WORKSPACE_TEMPLATE_CACHE_DIR", str(cache_root))

    _write_docx(
        external_root / "software_requirements_specification_machine.docx",
        [
            "[[SECTION:software_requirements_specification]]",
            "[[INSTRUCTION:software_requirements_specification_scope]]",
            "[[FIELD:software_scope]]",
            "[[VALUE:software_scope]]",
        ],
    )
    _write_docx(
        external_root / "software_requirements_specification_human.docx",
        [
            "[Document Title]",
            "[Optional Subtitle or Description]",
            "Document ID: [ID]",
            "Version: [Version]",
            "Effective Date: [Date]",
            "Prepared by: [Name / Company]",
            "Software Requirements Specification",
            "[Use this document to define the software requirements.]",
            "Software Scope",
            "[Describe the software scope.]",
        ],
    )

    spec = wt.discover_workspace_templates()[0]
    out_path = tmp_path / "rendered.docx"
    wt.render_workspace_template_to_docx(
        spec=spec,
        template_blocks={
            "instruction::software_requirements_specification_scope": "This document defines the release scope and boundary conditions.",
            "value::software_scope": "The software includes the clinician dashboard and data-ingestion services.",
        },
        output_path=str(out_path),
        document_metadata={
            "document_title": "Software Requirements Specification",
            "document_id": "SRS-001",
            "version": "0.3",
            "effective_date": "2026-03-19",
            "prepared_by": "NaviSsurance Test",
        },
    )

    rendered = Document(str(out_path))
    texts = [p.text for p in rendered.paragraphs]
    assert "Software Requirements Specification" in texts
    assert "Document ID: SRS-001" in texts
    assert "Version: 0.3" in texts
    assert "Effective Date: 2026-03-19" in texts
    assert "Prepared by: NaviSsurance Test" in texts
    assert "This document defines the release scope and boundary conditions." in texts
    assert "The software includes the clinician dashboard and data-ingestion services." in texts
    assert "[Describe the software scope.]" not in texts


def test_import_workspace_template_pair_copies_into_cache(tmp_path, monkeypatch):
    external_root = tmp_path / "external"
    cache_root = tmp_path / "cache"
    external_root.mkdir()
    cache_root.mkdir()
    monkeypatch.setattr(wt, "DEFAULT_WORKSPACE_TEMPLATE_ROOT", str(external_root))
    monkeypatch.setattr(wt, "WORKSPACE_TEMPLATE_CACHE_DIR", str(cache_root))

    machine = external_root / "user_needs_machine.docx"
    human = external_root / "user_needs_human.docx"
    _write_docx(machine, ["[[FIELD:primary_users]]", "[[VALUE:primary_users]]"])
    _write_docx(human, ["Primary Users", "[Describe the primary users.]"])

    spec = wt.import_workspace_template_pair(machine_path=str(machine), human_path=str(human))

    assert spec.source == "imported"
    imported_specs = [item for item in wt.discover_workspace_templates() if item.source == "imported"]
    assert imported_specs
    assert any(Path(item.machine_path).exists() and Path(item.human_path).exists() for item in imported_specs)


def test_import_workspace_template_pair_renames_nonstandard_files_for_discovery(tmp_path, monkeypatch):
    external_root = tmp_path / "external"
    cache_root = tmp_path / "cache"
    external_root.mkdir()
    cache_root.mkdir()
    monkeypatch.setattr(wt, "DEFAULT_WORKSPACE_TEMPLATE_ROOT", str(external_root))
    monkeypatch.setattr(wt, "WORKSPACE_TEMPLATE_CACHE_DIR", str(cache_root))

    machine = external_root / "URS Draft v7.docx"
    human = external_root / "URS Client Facing Template.dotx"
    _write_docx(machine, ["[[FIELD:primary_users]]", "[[VALUE:primary_users]]"])
    _write_docx(human, ["Primary Users", "[Describe the primary users.]"])

    spec = wt.import_workspace_template_pair(machine_path=str(machine), human_path=str(human))

    assert spec.source == "imported"
    imported_specs = [item for item in wt.discover_workspace_templates() if item.key == spec.key]
    assert len(imported_specs) == 1
    imported = imported_specs[0]
    assert Path(imported.machine_path).name.endswith("_machine.docx")
    assert Path(imported.human_path).name.endswith("_human.dotx")
    assert imported.values == ["primary_users"]


def test_build_template_contract_includes_explicit_schema(tmp_path, monkeypatch):
    external_root = tmp_path / "external"
    cache_root = tmp_path / "cache"
    external_root.mkdir()
    cache_root.mkdir()
    monkeypatch.setattr(wt, "DEFAULT_WORKSPACE_TEMPLATE_ROOT", str(external_root))
    monkeypatch.setattr(wt, "WORKSPACE_TEMPLATE_CACHE_DIR", str(cache_root))

    _write_docx(
        external_root / "plan_machine.docx",
        [
            "[[FIELD:purpose]]",
            "[[VALUE:purpose]]",
        ],
    )
    _write_docx(
        external_root / "plan_human.docx",
        [
            "Purpose",
            "[Describe the purpose.]",
        ],
    )

    spec = wt.discover_workspace_templates()[0]
    contract = wt.build_template_contract(spec)

    assert "Schema object to follow exactly:" in contract
    assert '"required_template_blocks"' in contract
    assert '"completion_analysis_keys"' in contract
    assert '"value::purpose"' in contract


def test_validate_template_payload_reports_missing_and_unexpected_keys(tmp_path, monkeypatch):
    external_root = tmp_path / "external"
    cache_root = tmp_path / "cache"
    external_root.mkdir()
    cache_root.mkdir()
    monkeypatch.setattr(wt, "DEFAULT_WORKSPACE_TEMPLATE_ROOT", str(external_root))
    monkeypatch.setattr(wt, "WORKSPACE_TEMPLATE_CACHE_DIR", str(cache_root))

    _write_docx(
        external_root / "policy_machine.docx",
        [
            "[[FIELD:scope]]",
            "[[VALUE:scope]]",
            "[[FIELD:responsibilities]]",
            "[[VALUE:responsibilities]]",
        ],
    )
    _write_docx(
        external_root / "policy_human.docx",
        [
            "Scope",
            "[Describe the scope.]",
            "Responsibilities",
            "[Describe the responsibilities.]",
        ],
    )

    spec = wt.discover_workspace_templates()[0]
    result = wt.validate_template_payload(
        spec,
        template_blocks={
            "value::scope": "Applies to released software.",
            "value::extra": "Unexpected",
        },
        document_metadata={
            "document_title": "Policy",
            "subtitle": "Draft",
            "document_id": "POL-001",
            "version": "0.1",
            "effective_date": "2026-03-19",
            "prepared_by": "Navi",
            "extra": "Unexpected",
        },
    )

    assert result.is_valid is False
    assert result.missing_block_keys == ["value::responsibilities"]
    assert result.unexpected_block_keys == ["value::extra"]
    assert result.unexpected_metadata_keys == ["extra"]
