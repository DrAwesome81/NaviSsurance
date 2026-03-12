from PyQt6.QtWidgets import (
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QTextEdit,
    QPushButton,
    QLineEdit,
    QLabel,
    QFileDialog,
    QListWidget,
    QListWidgetItem,
    QSplitter,
    QMessageBox,
)
from PyQt6.QtCore import Qt, QTimer, pyqtSignal, QThread
from datetime import datetime
import json
import re

from core.db import DatabaseManager
from docx import Document
from fpdf import FPDF


def robust_json_parse(response: str, logger=print):
    """
    Robustly parse JSON from an LLM response that *should* be JSON,
    but may have extra text or small formatting issues.
    Returns (data, success: bool).
    """
    import json
    import re

    if response is None:
        logger("DEBUG: robust_json_parse received None response")
        return None, False

    original_response = response
    response = response.strip()
    logger(f"DEBUG: Starting robust_json_parse. Raw response (truncated): {original_response[:200]}")

    try:
        data = json.loads(response)
        logger("DEBUG: Direct JSON parse succeeded")
        return data, True
    except json.JSONDecodeError as e:
        logger(f"DEBUG: Direct parse failed: {e}")

    json_match = re.search(r"\{.*\}", response, re.DOTALL)
    if json_match:
        candidate = json_match.group(0).strip()
        logger(f"DEBUG: Found JSON-looking block (truncated): {candidate[:200]}")
        try:
            data = json.loads(candidate)
            logger("DEBUG: Parsed JSON from extracted block")
            return data, True
        except json.JSONDecodeError as e:
            logger(f"DEBUG: Extracted block parse failed: {e}")

    cleaned = response.replace("\r", " ").replace("\n", " ")
    cleaned = re.sub(r"\s+", " ", cleaned)
    cleaned = re.sub(r",\s*}", "}", cleaned)
    cleaned = re.sub(r",\s*]", "]", cleaned)
    cleaned = cleaned.rstrip("'\"")
    logger(f"DEBUG: Attempting cleaned JSON parse (truncated): {cleaned[:200]}")

    try:
        data = json.loads(cleaned)
        logger("DEBUG: Cleaned JSON parse succeeded")
        return data, True
    except json.JSONDecodeError as e:
        logger(f"DEBUG: Cleaned parse still failed: {e}")

    cleaned_str = original_response.strip()
    if cleaned_str:
        logger("DEBUG: robust_json_parse failed to parse JSON; returning plain string fallback")
        return cleaned_str, True

    logger("DEBUG: robust_json_parse ultimately failed with empty/whitespace response")
    return None, False


class NoteProcessingThread(QThread):
    result_signal = pyqtSignal(str)
    error_signal = pyqtSignal(str)

    def __init__(self, chat_handler, prompt, session_id):
        super().__init__()
        self.chat_handler = chat_handler
        self.prompt = prompt
        self.session_id = session_id

    def run(self):
        try:
            response = self.chat_handler.get_response(self.prompt, self.session_id, conversation_history=[])
            self.result_signal.emit(response)
        except Exception as e:
            self.error_signal.emit(str(e))


class NoteTakingSystem(QWidget):
    def __init__(self, chat_handler):
        super().__init__()
        self.chat_handler = chat_handler
        self.db = DatabaseManager()
        self.db.init_notes_table()
        self.context = None
        self._current_note_id: int | None = None
        self._active_document_id: int | None = None
        self._loading_document = False
        self._document_dirty = False
        self._auto_select_initial_context = True
        self._reorganize_in_progress = False
        self.setup_ui()
        self.refresh_notes()

    def setup_ui(self):
        layout = QHBoxLayout(self)
        layout.setContentsMargins(8, 8, 8, 8)
        layout.setSpacing(8)

        input_style = (
            "background-color: #22252c; color: #e8eaed; border: 1px solid #2e2f32; "
            "padding: 8px 10px; border-radius: 6px;"
        )
        box_style = (
            "background-color: #22252c; color: #e8eaed; border: 1px solid #2e2f32; "
            "border-radius: 6px;"
        )
        btn_style = (
            "QPushButton { background-color: #FD6262; color: white; border: none; padding: 8px 12px; "
            "border-radius: 6px; font-weight: 500; }"
            "QPushButton:hover { background-color: rgba(253, 98, 98, 0.92); }"
            "QPushButton:disabled { background-color: #3a3d46; color: #9aa0a6; }"
        )

        splitter = QSplitter(Qt.Orientation.Horizontal)
        layout.addWidget(splitter)

        sidebar = QWidget()
        sidebar_layout = QVBoxLayout(sidebar)
        sidebar_layout.setContentsMargins(0, 0, 0, 0)
        sidebar_layout.setSpacing(6)
        sidebar_layout.addWidget(QLabel("Context Documents"))

        self.context_input = QLineEdit()
        self.context_input.setPlaceholderText("Set or create context (e.g., FDA protocol review)")
        self.context_input.setStyleSheet(input_style)
        self.context_input.returnPressed.connect(self.update_context)
        sidebar_layout.addWidget(self.context_input)

        sidebar_hint = QLabel("Each context is one living document.")
        sidebar_hint.setStyleSheet("color: #9aa0a6; font-size: 12px;")
        sidebar_layout.addWidget(sidebar_hint)

        self.context_list = QListWidget()
        self.context_list.setStyleSheet(box_style)
        self.context_list.itemSelectionChanged.connect(self._on_context_selected)
        sidebar_layout.addWidget(self.context_list, 1)

        splitter.addWidget(sidebar)

        main = QWidget()
        main_layout = QVBoxLayout(main)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(6)

        self.active_context_label = QLabel("Current document: (none)")
        self.active_context_label.setStyleSheet("color: #e8eaed; font-weight: 600; font-size: 13px;")
        main_layout.addWidget(self.active_context_label)

        self.doc_meta_label = QLabel("")
        self.doc_meta_label.setStyleSheet("color: #9aa0a6; font-size: 12px;")
        main_layout.addWidget(self.doc_meta_label)

        capture_label = QLabel("Add observation:")
        capture_label.setStyleSheet("color: #e8eaed; font-weight: 600; font-size: 13px;")
        main_layout.addWidget(capture_label)

        class CaptureTextEdit(QTextEdit):
            def __init__(self, parent=None):
                super().__init__(parent)
                self.parent_widget = parent

            def keyPressEvent(self, event):
                if event.key() == Qt.Key.Key_Return and not event.modifiers() & Qt.KeyboardModifier.ShiftModifier:
                    if hasattr(self.parent_widget, "process_note"):
                        self.parent_widget.process_note()
                    event.accept()
                    return
                super().keyPressEvent(event)

        self.chat_input = CaptureTextEdit(self)
        self.chat_input.setPlaceholderText("Type notes while reviewing. Press Enter to merge into the current document.")
        self.chat_input.setStyleSheet(box_style)
        self.chat_input.setFixedHeight(110)
        main_layout.addWidget(self.chat_input)

        self.note_status_label = QLabel("")
        self.note_status_label.setStyleSheet("color: #9aa0a6; font-size: 11px;")
        self.note_status_label.setVisible(False)
        main_layout.addWidget(self.note_status_label)

        self.document_editor = QTextEdit()
        self.document_editor.setStyleSheet(box_style)
        self.document_editor.setPlaceholderText(
            "Select or create a context, then add observations.\n\n"
            "Navi will maintain one structured document for that context."
        )
        self.document_editor.textChanged.connect(self._on_document_text_changed)
        main_layout.addWidget(self.document_editor, 1)

        actions = QHBoxLayout()
        self.save_btn = QPushButton("Save Document")
        self.save_btn.setStyleSheet(btn_style)
        self.save_btn.clicked.connect(self.save_current_document)
        actions.addWidget(self.save_btn)

        self.copy_btn = QPushButton("Copy")
        self.copy_btn.setStyleSheet(btn_style)
        self.copy_btn.clicked.connect(self.copy_current_document)
        actions.addWidget(self.copy_btn)

        self.reorg_btn = QPushButton("Re-organize")
        self.reorg_btn.setStyleSheet(btn_style)
        self.reorg_btn.clicked.connect(self.reorganize_notes)
        actions.addWidget(self.reorg_btn)

        self.export_btn = QPushButton("Export…")
        self.export_btn.setStyleSheet(btn_style)
        self.export_btn.clicked.connect(self.export_notes)
        actions.addWidget(self.export_btn)
        actions.addStretch()
        main_layout.addLayout(actions)

        splitter.addWidget(main)
        splitter.setStretchFactor(0, 1)
        splitter.setStretchFactor(1, 3)
        self.setLayout(layout)
        self._sync_action_states()

    def update_context(self):
        value = (self.context_input.text() or "").strip()
        if not value:
            self.context = None
            self._active_document_id = None
            self._auto_select_initial_context = False
            self.context_list.clearSelection()
            self.context_input.clear()
            self.refresh_notes()
            return

        doc = self.db.get_note_document(value, create=True)
        self.context = (doc or {}).get("context") or value
        self._auto_select_initial_context = True
        self.context_input.clear()
        self.refresh_notes()

    def _on_context_selected(self):
        items = self.context_list.selectedItems()
        if not items:
            return
        self.context = (items[0].data(Qt.ItemDataRole.UserRole) or "").strip() or None
        self._auto_select_initial_context = True
        self.refresh_notes()

    def _on_document_text_changed(self):
        if self._loading_document:
            return
        self._document_dirty = True
        self._sync_action_states()
        if self.context:
            self.note_status_label.setText("Document has unsaved manual edits.")
            self.note_status_label.setVisible(True)

    @staticmethod
    def _note_ts_in_days(note_ts: str, days: int) -> bool:
        try:
            dt = datetime.strptime((note_ts or "").strip(), "%Y-%m-%d %H:%M:%S")
            return (datetime.now() - dt).days <= int(days)
        except Exception:
            return True

    def refresh_notes(self):
        try:
            documents = self.db.list_note_documents(limit=500)
            contexts = [doc.get("context") or "" for doc in documents if (doc.get("context") or "").strip()]

            if self.context is None and contexts and self._auto_select_initial_context:
                self.context = contexts[0]

            self.context_list.blockSignals(True)
            self.context_list.clear()
            for doc in documents:
                context = (doc.get("context") or "").strip()
                if not context:
                    continue
                count = int(doc.get("source_note_count") or 0)
                label = f"{context} ({count})" if count else context
                item = QListWidgetItem(label)
                item.setData(Qt.ItemDataRole.UserRole, context)
                self.context_list.addItem(item)
                if self.context == context:
                    self.context_list.setCurrentItem(item)
            self.context_list.blockSignals(False)

            doc = self.db.get_note_document(self.context, create=False) if self.context else None
            self._active_document_id = int(doc.get("id") or 0) if doc else None
            self._load_document(doc)
        except Exception as e:
            self._loading_document = True
            self.document_editor.setPlainText(f"Error loading notes document: {e}")
            self._loading_document = False
            self.note_status_label.setText(f"Error loading notes document: {e}")
            self.note_status_label.setVisible(True)
            self._sync_action_states()

    def _load_document(self, doc: dict | None):
        self._loading_document = True
        try:
            if not doc:
                self.active_context_label.setText("Current document: (none)")
                self.doc_meta_label.setText("Set a context to start building a document.")
                self.document_editor.setPlainText("")
                self.note_status_label.setVisible(False)
                self.note_status_label.setText("")
                self._document_dirty = False
                self._sync_action_states()
                return

            context = (doc.get("context") or "").strip()
            state = (doc.get("state") or "ready").strip().lower()
            error_text = (doc.get("error_text") or "").strip()
            count = int(doc.get("source_note_count") or 0)
            updated_at = (doc.get("updated_at") or "").strip()
            self.active_context_label.setText(f"Current document: {context}")
            meta_bits = [f"Source captures: {count}"]
            if updated_at:
                meta_bits.append(f"Updated: {updated_at}")
            self.doc_meta_label.setText(" | ".join(meta_bits))
            self.document_editor.setPlainText((doc.get("document_text") or "").strip())
            self._document_dirty = False

            if state == "pending":
                self.note_status_label.setText("Updating document…")
                self.note_status_label.setVisible(True)
            elif state == "error":
                self.note_status_label.setText(error_text or "Document update failed.")
                self.note_status_label.setVisible(True)
            else:
                self.note_status_label.setVisible(False)
                self.note_status_label.setText("")
            self._sync_action_states()
        finally:
            self._loading_document = False

    def _sync_action_states(self):
        has_context = bool((self.context or "").strip())
        has_doc_text = bool((self.document_editor.toPlainText() or "").strip())
        self.chat_input.setEnabled(has_context)
        self.document_editor.setEnabled(has_context)
        self.save_btn.setEnabled(has_context and self._document_dirty)
        self.copy_btn.setEnabled(has_doc_text)
        self.reorg_btn.setEnabled(has_context and has_doc_text and not self._document_dirty)
        self.export_btn.setEnabled(has_context and has_doc_text)

    def process_note(self):
        if not self.context:
            QMessageBox.information(self, "Notes", "Set a context before adding observations.")
            return

        content = (self.chat_input.toPlainText() or "").strip()
        if not content:
            return

        ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        note_id = self.db.create_note_draft(raw_note=content, timestamp=ts, context=self.context, category="Captured")
        self._current_note_id = note_id or None
        existing_doc = self.db.get_note_document(self.context, create=True) or {}
        self.db.save_note_document(
            self.context,
            document_text=(existing_doc.get("document_text") or "").strip(),
            title=self.context,
            state="pending",
            error_text=None,
        )
        self.note_status_label.setText("Updating document…")
        self.note_status_label.setVisible(True)
        self.chat_input.setEnabled(False)
        self._sync_action_states()
        self._start_document_update_for_note(
            note_id=int(note_id),
            raw_text=content,
            context=self.context,
            current_document=(existing_doc.get("document_text") or "").strip(),
            retry_count=0,
        )

    def _start_document_update_for_note(
        self,
        *,
        note_id: int,
        raw_text: str,
        context: str,
        current_document: str,
        retry_count: int = 0,
    ) -> None:
        prompt = f"""Context: {context}

You maintain a single working notes document for this context.

CURRENT DOCUMENT:
<<<DOCUMENT
{current_document or "(empty)"}
DOCUMENT>>>

NEW OBSERVATION:
{raw_text}

INSTRUCTIONS:
- Merge the new observation into the current document.
- Rewrite for clarity, but preserve the meaning.
- Organize the document into useful section headings and concise bullet points.
- Consolidate duplicate or overlapping bullets instead of repeating them.
- Keep the whole document coherent as if it were one evolving work product.
- If the document is empty, create a sensible structure.

RESPONSE FORMAT:
Return ONLY valid JSON in this exact shape:
{{"document": "Updated document text here"}}
"""
        self.note_thread = NoteProcessingThread(self.chat_handler, prompt, "notes_session")
        self.note_thread.result_signal.connect(
            lambda resp, nid=note_id, raw=raw_text, ctx=context, current=current_document, retries=retry_count:
                self.handle_note_formatted(resp, nid, raw, ctx, current, retries)
        )
        self.note_thread.error_signal.connect(lambda error, nid=note_id, ctx=context: self.handle_note_error(error, nid, ctx))
        self.note_thread.start()

    def handle_note_formatted(self, response, note_id: int, raw_text: str, context: str, current_document: str, retry_count: int):
        QTimer.singleShot(
            0,
            lambda: self._handle_note_formatted_safe(response, note_id, raw_text, context, current_document, retry_count),
        )

    def _handle_note_formatted_safe(
        self,
        response,
        note_id: int,
        raw_text: str,
        context: str,
        current_document: str = "",
        retry_count: int = 0,
    ):
        try:
            if self._looks_like_worker_failure(str(response or "")) and retry_count < 1:
                self.note_status_label.setText("Local AI hiccuped. Retrying once…")
                self.note_status_label.setVisible(True)
                QTimer.singleShot(
                    300,
                    lambda: self._start_document_update_for_note(
                        note_id=note_id,
                        raw_text=raw_text,
                        context=context,
                        current_document=current_document,
                        retry_count=retry_count + 1,
                    ),
                )
                return
            parsed, success = robust_json_parse(response, logger=lambda msg: print(f"Note formatting: {msg}"))
            document_text = self._extract_document_text(parsed, response)
            if success and document_text:
                self.db.update_note_by_id(
                    int(note_id),
                    formatted_note=raw_text.strip(),
                    category="Merged",
                    context=context,
                    state="ready",
                    error_text=None,
                )
                self.db.save_note_document(
                    context,
                    document_text=document_text.strip(),
                    title=context,
                    state="ready",
                    error_text=None,
                    source_note_count=self.db.count_ready_notes_for_context(context),
                )
                self.chat_input.clear()
                self.chat_input.setEnabled(True)
                self.chat_input.setFocus()
                self.refresh_notes()
                return

            self.db.update_note_by_id(int(note_id), state="error", error_text="Could not parse document update")
            self.db.save_note_document(
                context,
                document_text=(self.db.get_note_document(context, create=True) or {}).get("document_text", ""),
                title=context,
                state="error",
                error_text="Could not parse document update",
            )
            QMessageBox.warning(self, "Notes", "Could not parse the updated document. The raw observation was kept in history.")
            self.chat_input.setEnabled(True)
            self.refresh_notes()
        except Exception as e:
            self.db.update_note_by_id(int(note_id), state="error", error_text=str(e))
            self.db.save_note_document(
                context,
                document_text=(self.db.get_note_document(context, create=True) or {}).get("document_text", ""),
                title=context,
                state="error",
                error_text=str(e),
            )
            QMessageBox.warning(self, "Notes", f"Error processing note: {e}")
            self.chat_input.setEnabled(True)
            self.refresh_notes()

    def handle_note_error(self, error, note_id: int | None = None, context: str | None = None):
        QTimer.singleShot(0, lambda: self._handle_note_error_safe(error, note_id, context))

    def _handle_note_error_safe(self, error, note_id: int | None = None, context: str | None = None):
        if note_id:
            self.db.update_note_by_id(int(note_id), state="error", error_text=str(error))
        if context:
            existing_doc = self.db.get_note_document(context, create=True) or {}
            self.db.save_note_document(
                context,
                document_text=(existing_doc.get("document_text") or "").strip(),
                title=context,
                state="error",
                error_text=str(error),
                source_note_count=self.db.count_ready_notes_for_context(context),
            )
        QMessageBox.warning(self, "Notes", f"Error processing note: {error}")
        self.chat_input.setEnabled(True)
        self.refresh_notes()

    def reorganize_notes(self, *, retry_count: int = 0):
        if not self.context:
            return
        if self._reorganize_in_progress:
            return
        current_document = (self.document_editor.toPlainText() or "").strip()
        if not current_document:
            return
        self._reorganize_in_progress = True
        self.db.save_note_document(
            self.context,
            document_text=current_document,
            title=self.context,
            state="pending",
            error_text=None,
            source_note_count=self.db.count_ready_notes_for_context(self.context),
        )
        self.note_status_label.setText("Re-organizing document…")
        self.note_status_label.setVisible(True)
        self.reorg_btn.setEnabled(False)
        prompt = f"""Context: {self.context}

You are refining a working notes document for this context.

DOCUMENT:
<<<DOCUMENT
{current_document}
DOCUMENT>>>

INSTRUCTIONS:
- Reorganize the document into cleaner sections and tighter bullet points.
- Preserve the substance of the document.
- Merge duplicates and improve clarity.
- Keep it concise, but do not drop important details.

RESPONSE FORMAT:
Return ONLY valid JSON:
{{"document": "Refined document text here"}}
"""
        self.organize_thread = NoteProcessingThread(self.chat_handler, prompt, "notes_session")
        self.organize_thread.result_signal.connect(
            lambda resp, ctx=self.context, current=current_document, retries=retry_count:
                self.handle_reorganization_result(resp, ctx, current, retries)
        )
        self.organize_thread.error_signal.connect(lambda error, ctx=self.context: self.handle_reorganization_error(error, ctx))
        self.organize_thread.start()

    def handle_reorganization_result(self, response, context: str, current_document: str, retry_count: int):
        QTimer.singleShot(0, lambda: self._handle_reorganization_result_safe(response, context, current_document, retry_count))

    def _handle_reorganization_result_safe(self, response, context: str, current_document: str = "", retry_count: int = 0):
        try:
            if self._looks_like_worker_failure(str(response or "")) and retry_count < 1:
                self.note_status_label.setText("Local AI hiccuped. Retrying re-organization once…")
                self.note_status_label.setVisible(True)
                self._reorganize_in_progress = False
                QTimer.singleShot(300, lambda: self.reorganize_notes(retry_count=retry_count + 1))
                return
            parsed, success = robust_json_parse(response, logger=print)
            document_text = self._extract_document_text(parsed, response)
            if success and document_text:
                self.db.save_note_document(
                    context,
                    document_text=document_text.strip(),
                    title=context,
                    state="ready",
                    error_text=None,
                    source_note_count=self.db.count_ready_notes_for_context(context),
                )
                self.refresh_notes()
                return
            self.db.save_note_document(
                context,
                    document_text=current_document or (self.document_editor.toPlainText() or "").strip(),
                title=context,
                state="error",
                error_text="Could not parse reorganized document",
                source_note_count=self.db.count_ready_notes_for_context(context),
            )
            QMessageBox.warning(self, "Notes", "Could not parse the reorganized document.")
            self.refresh_notes()
        except Exception as e:
            self.db.save_note_document(
                context,
                document_text=current_document or (self.document_editor.toPlainText() or "").strip(),
                title=context,
                state="error",
                error_text=str(e),
                source_note_count=self.db.count_ready_notes_for_context(context),
            )
            QMessageBox.warning(self, "Notes", f"Error reorganizing document: {e}")
            self.refresh_notes()
        finally:
            self._reorganize_in_progress = False

    def handle_reorganization_error(self, error, context: str):
        QTimer.singleShot(0, lambda: self._handle_reorganization_error_safe(error, context))

    def _handle_reorganization_error_safe(self, error, context: str):
        self.db.save_note_document(
            context,
            document_text=(self.document_editor.toPlainText() or "").strip(),
            title=context,
            state="error",
            error_text=str(error),
            source_note_count=self.db.count_ready_notes_for_context(context),
        )
        QMessageBox.warning(self, "Notes", f"Error reorganizing document: {error}")
        self.refresh_notes()
        self._reorganize_in_progress = False

    @staticmethod
    def _extract_document_text(parsed, response: str) -> str | None:
        if isinstance(parsed, dict):
            for key in ("document", "document_text", "updated_document", "formatted", "content"):
                value = parsed.get(key)
                if isinstance(value, str) and value.strip():
                    candidate = value.strip()
                    if not NoteTakingSystem._looks_like_worker_failure(candidate):
                        return candidate
        if isinstance(parsed, list):
            for item in parsed:
                if isinstance(item, dict):
                    candidate = NoteTakingSystem._extract_document_text(item, response)
                    if candidate:
                        return candidate
                elif isinstance(item, str) and item.strip():
                    candidate = item.strip()
                    if not NoteTakingSystem._looks_like_worker_failure(candidate):
                        return candidate
        if isinstance(parsed, str) and parsed.strip():
            candidate = parsed.strip()
            if not NoteTakingSystem._looks_like_worker_failure(candidate):
                return candidate
        if '"document"' in response:
            match = re.search(r'"document"\s*:\s*"(.+?)"', response, re.DOTALL)
            if match:
                candidate = match.group(1).strip()
                if not NoteTakingSystem._looks_like_worker_failure(candidate):
                    return candidate
        return None

    @staticmethod
    def _looks_like_worker_failure(text: str) -> bool:
        lowered = (text or "").strip().lower()
        if not lowered:
            return True
        return lowered.startswith("error:") or lowered in {
            "model not loaded.",
            "local ai's acting up—try again.",
            "local ai's acting up-try again.",
            "local ai's acting up, try again.",
        } or "worker communication error" in lowered or "local ai's acting up" in lowered

    def save_current_document(self):
        if not self.context:
            return
        self.db.save_note_document(
            self.context,
            document_text=(self.document_editor.toPlainText() or "").strip(),
            title=self.context,
            state="ready",
            error_text=None,
            source_note_count=self.db.count_ready_notes_for_context(self.context),
        )
        self._document_dirty = False
        self.note_status_label.setVisible(False)
        self.note_status_label.setText("")
        self.refresh_notes()

    def copy_current_document(self):
        text = (self.document_editor.toPlainText() or "").strip()
        if not text:
            return
        from PyQt6.QtWidgets import QApplication

        QApplication.clipboard().setText(text)
        QMessageBox.information(self, "Notes", "Document copied to clipboard.")

    @staticmethod
    def _slugify_filename(value: str) -> str:
        slug = re.sub(r"[^A-Za-z0-9._-]+", "_", (value or "").strip())
        return slug.strip("._") or "notes_document"

    def export_notes(self):
        if not self.context:
            QMessageBox.information(self, "Export", "Select a context document first.")
            return
        text = (self.document_editor.toPlainText() or "").strip()
        if not text:
            QMessageBox.information(self, "Export", "No document content to export.")
            return

        default_filename = f"{self._slugify_filename(self.context)}.docx"
        file_path, selected_filter = QFileDialog.getSaveFileName(
            self,
            "Export Notes Document",
            default_filename,
            "Word Document (*.docx);;Markdown (*.md);;PDF (*.pdf)",
        )
        if not file_path:
            return

        ext = file_path.lower().rsplit(".", 1)[-1] if "." in file_path else ""
        if selected_filter.startswith("Word") or ext == "docx":
            if not file_path.lower().endswith(".docx"):
                file_path += ".docx"
            doc = Document()
            doc.add_heading(self.context, level=1)
            doc.add_paragraph(f"Exported: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            doc.add_paragraph("")
            self._write_docx_document(doc, text)
            doc.save(file_path)
            QMessageBox.information(self, "Export", f"Exported to {file_path}")
            return

        if selected_filter.startswith("Markdown") or ext == "md":
            if not file_path.lower().endswith(".md"):
                file_path += ".md"
            lines = [
                f"# {self.context}",
                f"**Exported:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
                "",
                text,
                "",
            ]
            with open(file_path, "w", encoding="utf-8") as f:
                f.write("\n".join(lines))
            QMessageBox.information(self, "Export", f"Exported to {file_path}")
            return

        if not file_path.lower().endswith(".pdf"):
            file_path += ".pdf"
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        pdf.set_font("Helvetica", style="B", size=14)
        pdf.multi_cell(0, 8, self.context)
        pdf.ln(1)
        pdf.set_font("Helvetica", size=10)
        pdf.multi_cell(0, 6, f"Exported: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        pdf.ln(2)
        self._write_pdf_document(pdf, text)
        pdf.output(file_path)
        QMessageBox.information(self, "Export", f"Exported to {file_path}")

    @staticmethod
    def _write_docx_document(doc: Document, text: str) -> None:
        for line in text.splitlines():
            stripped = line.strip()
            if not stripped:
                doc.add_paragraph("")
            elif stripped.startswith("### "):
                doc.add_heading(stripped[4:].strip(), level=3)
            elif stripped.startswith("## "):
                doc.add_heading(stripped[3:].strip(), level=2)
            elif stripped.startswith("# "):
                doc.add_heading(stripped[2:].strip(), level=1)
            elif stripped.startswith("- ") or stripped.startswith("* "):
                doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
            else:
                doc.add_paragraph(stripped)

    @staticmethod
    def _write_pdf_document(pdf: FPDF, text: str) -> None:
        for line in text.splitlines():
            stripped = line.strip()
            if not stripped:
                pdf.ln(2)
            elif stripped.startswith("### "):
                pdf.set_font("Helvetica", style="B", size=11)
                pdf.multi_cell(0, 6, stripped[4:].strip())
                pdf.set_font("Helvetica", size=10)
            elif stripped.startswith("## "):
                pdf.set_font("Helvetica", style="B", size=12)
                pdf.multi_cell(0, 7, stripped[3:].strip())
                pdf.set_font("Helvetica", size=10)
            elif stripped.startswith("# "):
                pdf.set_font("Helvetica", style="B", size=13)
                pdf.multi_cell(0, 8, stripped[2:].strip())
                pdf.set_font("Helvetica", size=10)
            elif stripped.startswith("- ") or stripped.startswith("* "):
                pdf.multi_cell(0, 6, f"- {stripped[2:].strip()}")
            else:
                pdf.multi_cell(0, 6, stripped)
