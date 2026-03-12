from __future__ import annotations

from datetime import datetime
import os

from docx import Document
from fpdf import FPDF


def export_markdownish_document(
    *,
    title: str,
    text: str,
    file_path: str,
    selected_filter: str = "",
    exported_at: datetime | None = None,
) -> str:
    body = (text or "").strip()
    if not body:
        raise ValueError("No document content to export.")

    export_time = exported_at or datetime.now()
    export_stamp = export_time.strftime("%Y-%m-%d %H:%M:%S")
    selected = (selected_filter or "").lower()
    ext = os.path.splitext(str(file_path or ""))[1].lower()

    if "docx" in selected or ext == ".docx":
        if ext != ".docx":
            file_path += ".docx"
        doc = Document()
        doc.add_heading(title or "Document", level=1)
        doc.add_paragraph(f"Exported: {export_stamp}")
        doc.add_paragraph("")
        _write_docx_document(doc, body)
        doc.save(file_path)
        return file_path

    if "pdf" in selected or ext == ".pdf":
        if ext != ".pdf":
            file_path += ".pdf"
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        pdf.set_font("Helvetica", style="B", size=14)
        pdf.multi_cell(0, 8, title or "Document")
        pdf.ln(1)
        pdf.set_font("Helvetica", size=10)
        pdf.multi_cell(0, 6, f"Exported: {export_stamp}")
        pdf.ln(2)
        _write_pdf_document(pdf, body)
        pdf.output(file_path)
        return file_path

    if "markdown" in selected or ext == ".md":
        if ext != ".md":
            file_path += ".md"
        with open(file_path, "w", encoding="utf-8") as handle:
            handle.write(
                "\n".join(
                    [
                        f"# {title or 'Document'}",
                        f"**Exported:** {export_stamp}",
                        "",
                        body,
                        "",
                    ]
                )
            )
        return file_path

    if ext != ".txt":
        file_path += ".txt"
    with open(file_path, "w", encoding="utf-8") as handle:
        handle.write(body)
    return file_path


def _write_docx_document(doc: Document, text: str) -> None:
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph("")
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=1)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(stripped)


def _write_pdf_document(pdf: FPDF, text: str) -> None:
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            pdf.ln(2)
        elif stripped.startswith("### "):
            pdf.set_font("Helvetica", style="B", size=11)
            pdf.multi_cell(0, 6, stripped[4:].strip())
            pdf.set_font("Helvetica", size=10)
        elif stripped.startswith("## "):
            pdf.set_font("Helvetica", style="B", size=12)
            pdf.multi_cell(0, 7, stripped[3:].strip())
            pdf.set_font("Helvetica", size=10)
        elif stripped.startswith("# "):
            pdf.set_font("Helvetica", style="B", size=13)
            pdf.multi_cell(0, 8, stripped[2:].strip())
            pdf.set_font("Helvetica", size=10)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            pdf.multi_cell(0, 6, f"- {stripped[2:].strip()}")
        else:
            pdf.multi_cell(0, 6, stripped)
